# -*- coding: utf-8 -*-
"""对 word-deliverable 生成的 DOCX 做可测试、spec 感知的结构验收。"""
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

PROCESS_WORDS = [
    "请按 F9", "请按F9", "请更新域", "请右键更新", "请更新目录",
    "TODO", "待生成", "待排版", "示例占位符", "AI 说明", "AI说明",
    "请用户自行", "目录将在打开后自动生成", "请用户自行刷新",
    "[图片加载失败", "[图片缺失", "图片加载失败:", "图片缺失:",
]
HEADING_STYLES = [f"Heading {level}" for level in range(1, 7)]
REQUIRED_STYLES = HEADING_STYLES + ["小节标题", "表题", "图题"]


def _part_roots(doc):
    """返回正文以及去重后的页眉、页脚 XML 根节点。"""
    roots = [("document", doc.element)]
    seen = {id(doc.element)}
    for index, section in enumerate(doc.sections):
        candidates = (
            (f"section{index + 1}.header", section.header._element),
            (f"section{index + 1}.footer", section.footer._element),
            (f"section{index + 1}.firstHeader", section.first_page_header._element),
            (f"section{index + 1}.firstFooter", section.first_page_footer._element),
        )
        for name, root in candidates:
            if id(root) not in seen:
                roots.append((name, root))
                seen.add(id(root))
    return roots


def iter_complex_fields(root):
    """解析复杂域，验证 begin → instrText → separate → end 的完整顺序。"""
    active = []
    for element in root.iter():
        if element.tag == qn("w:fldChar"):
            field_type = element.get(qn("w:fldCharType"))
            if field_type == "begin":
                active.append({"instruction": "", "has_separate": False, "has_end": False})
            elif field_type == "separate" and active:
                active[-1]["has_separate"] = True
            elif field_type == "end" and active:
                field = active.pop()
                field["has_end"] = True
                field["complete"] = bool(field["instruction"].strip() and field["has_separate"])
                yield field
        elif element.tag == qn("w:instrText") and active:
            active[-1]["instruction"] += element.text or ""
    for field in active:
        field["complete"] = False
        yield field


def _field_name(instruction):
    match = re.match(r"\s*([A-Za-z]+)", instruction or "")
    return match.group(1).upper() if match else ""


def _check_fields(doc, spec, report, problems):
    fields = []
    for part_name, root in _part_roots(doc):
        for field in iter_complex_fields(root):
            field["part"] = part_name
            field["name"] = _field_name(field["instruction"])
            fields.append(field)

    required = []
    if spec is None or spec.get("toc", True) is True:
        required.append("TOC")
    if spec is None or spec.get("header_footer", True) is True:
        required.extend(["PAGE", "NUMPAGES"])

    if not required:
        report.append("[字段] TOC 与页眉页脚均按 spec 跳过")
    else:
        summary = []
        for name in required:
            matches = [field for field in fields if field["name"] == name]
            summary.append(f"{name}={bool(matches)}")
            if not matches:
                problems.append(f"缺少 {name} 字段")
            for field in matches:
                if not field["complete"]:
                    problems.append(f"{name} 域结构不完整（{field['part']}，应为 begin→instrText→separate→end）")
        report.append("[字段] " + " ".join(summary))

    toc_required = spec is None or spec.get("toc", True) is True
    placeholder = any("目录将在打开后自动生成" in paragraph.text for paragraph in doc.paragraphs)
    if toc_required:
        report.append(f"[目录] 已展开缓存={not placeholder}")
        if placeholder:
            problems.append("TOC 仍含占位文本，需运行 update_fields.py")
    else:
        report.append("[目录] 按 spec 跳过")


def _numbering_ids(doc):
    try:
        root = doc.part.numbering_part.element
    except Exception:
        return set()
    return {
        element.get(qn("w:numId"))
        for element in root.findall(qn("w:num"))
        if element.get(qn("w:numId")) is not None
    }


def _check_numbering(doc, report, problems):
    valid_num_ids = _numbering_ids(doc)
    bound = 0
    for level, style_name in enumerate(HEADING_STYLES):
        style = doc.styles[style_name]
        p_pr = style.element.find(qn("w:pPr"))
        num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
        num_id_element = num_pr.find(qn("w:numId")) if num_pr is not None else None
        ilvl_element = num_pr.find(qn("w:ilvl")) if num_pr is not None else None
        num_id = num_id_element.get(qn("w:val")) if num_id_element is not None else None
        ilvl = ilvl_element.get(qn("w:val")) if ilvl_element is not None else None
        if num_id is None or ilvl is None:
            problems.append(f"{style_name} 样式未完整绑定自动编号（缺少 numId/ilvl）")
            continue
        if num_id not in valid_num_ids:
            problems.append(f"{style_name} 编号 numId={num_id} 在 numbering.xml 中无定义")
            continue
        if ilvl != str(level):
            problems.append(f"{style_name} 编号层级 ilvl={ilvl}，预期 {level}")
            continue
        bound += 1
    report.append(f"[编号] H1-H6 完整绑定={bound}/6")


def _row_has(row, tag):
    tr_pr = row._tr.find(qn("w:trPr"))
    return tr_pr is not None and tr_pr.find(qn(tag)) is not None


def _check_tables(doc, spec, report, problems):
    fixed_count = 0
    cover_table_index = 0 if spec is None or spec.get("cover", True) is True else None
    for index, table in enumerate(doc.tables):
        tbl_pr = table._tbl.tblPr
        layout = tbl_pr.find(qn("w:tblLayout"))
        fixed = layout is not None and layout.get(qn("w:type")) == "fixed"
        fixed_count += int(fixed)
        if not fixed:
            problems.append(f"表格#{index} 未设置固定布局 tblLayout=fixed")
        if not table.rows:
            problems.append(f"表格#{index} 无数据行")
            continue
        if len(table.rows) > 1 and index != cover_table_index and not _row_has(table.rows[0], "w:tblHeader"):
            problems.append(f"表格#{index} 无重复标题行（行数={len(table.rows)}）")
        for row_index, row in enumerate(table.rows):
            if not _row_has(row, "w:cantSplit"):
                problems.append(f"表格#{index} 第{row_index + 1}行未禁止跨页拆分")
        grid = table._tbl.tblGrid.findall(qn("w:gridCol"))
        if len(grid) != len(table.columns):
            problems.append(f"表格#{index} 网格列数 {len(grid)} 与实际列数 {len(table.columns)} 不一致")
        widths = [int(col.get(qn("w:w"), "0")) for col in grid]
        if any(width <= 0 for width in widths):
            problems.append(f"表格#{index} 存在非正数网格列宽")
    report.append(f"[表格] 共 {len(doc.tables)} 个；固定布局 {fixed_count}/{len(doc.tables)}")


def _check_sections(doc, spec, report, problems):
    sections = doc.sections
    report.append(f"[节] 共 {len(sections)} 个 section")
    cover_enabled = spec is None or spec.get("cover", True) is True
    header_footer_enabled = spec is None or spec.get("header_footer", True) is True
    for index, section in enumerate(sections):
        landscape = section.orientation == WD_ORIENT.LANDSCAPE
        expected_width = 29.7 if landscape else 21.0
        expected_height = 21.0 if landscape else 29.7
        if abs(section.page_width.cm - expected_width) > 0.15 or abs(section.page_height.cm - expected_height) > 0.15:
            problems.append(f"Section#{index + 1} 页面尺寸/方向异常")
        if min(section.top_margin.cm, section.bottom_margin.cm, section.left_margin.cm, section.right_margin.cm) <= 0:
            problems.append(f"Section#{index + 1} 存在非正页边距")

    first = sections[0]
    if cover_enabled and header_footer_enabled:
        if not first.different_first_page_header_footer:
            problems.append("启用封面和页眉页脚时未设置首页不同")
        if first.first_page_header.paragraphs[0].text.strip():
            problems.append("封面页眉非空")
        if first.first_page_footer.paragraphs[0].text.strip():
            problems.append("封面页脚非空")
    elif first.different_first_page_header_footer:
        problems.append("未同时启用封面和页眉页脚，却设置了首页不同")

    if header_footer_enabled:
        if not first.header.paragraphs[0].text.strip():
            problems.append("正文页眉为空")
        for index, section in enumerate(sections[1:], start=2):
            if not section.header.is_linked_to_previous or not section.footer.is_linked_to_previous:
                problems.append(f"Section#{index} 页眉页脚未链接前节")
    else:
        for index, section in enumerate(sections, start=1):
            if section.header.paragraphs[0].text.strip() or section.footer.paragraphs[0].text.strip():
                problems.append(f"header_footer=false 但 Section#{index} 仍有页眉页脚内容")


def _check_text_and_headings(doc, spec, report, problems):
    found = []
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for word in PROCESS_WORDS:
            if word in paragraph.text:
                found.append((word, paragraph.text[:40]))
    for word, text in found:
        problems.append(f"过程文字 '{word}' in '{text}'")
    report.append(f"[过程文字] 发现 {len(found)} 处")

    sequence = []
    counts = {name: 0 for name in HEADING_STYLES}
    for paragraph in doc.paragraphs:
        if paragraph.style.name in counts:
            counts[paragraph.style.name] += 1
            sequence.append((int(paragraph.style.name.split()[-1]), paragraph.text.strip()))
    report.append("[标题] " + " ".join(f"H{i}={counts[f'Heading {i}']}" for i in range(1, 7)))
    if counts["Heading 1"] == 0:
        problems.append("未发现一级标题")
    for previous, current in zip(sequence, sequence[1:]):
        if current[0] > previous[0] + 1:
            problems.append(f"标题跳级: H{previous[0]}「{previous[1][:20]}」→ H{current[0]}「{current[1][:20]}」")
    manual_number = re.compile(r"^(\d+\.)+\d*\s")
    for level, text in sequence:
        if manual_number.match(text):
            problems.append(f"手工编号残留（应使用自动编号）: H{level}「{text[:30]}」")
    if spec and spec.get("document_title"):
        title = str(spec["document_title"]).strip()
        if any(level == 1 and text == title for level, text in sequence):
            problems.append(f"文档标题「{title}」被误编为第一章")


def validate_docx(path, spec=None):
    """返回 (report, problems)，不退出进程，供 CLI 与测试复用。"""
    doc = Document(str(path))
    report, problems = [], []
    style_names = {style.name for style in doc.styles}
    missing = [name for name in REQUIRED_STYLES if name not in style_names]
    problems.extend(f"缺失样式: {name}" for name in missing)
    report.append(f"[样式] 必需样式存在: {not missing}")
    _check_text_and_headings(doc, spec, report, problems)
    _check_numbering(doc, report, problems)
    _check_fields(doc, spec, report, problems)
    _check_tables(doc, spec, report, problems)
    _check_sections(doc, spec, report, problems)
    return report, problems


def main():
    if len(sys.argv) < 2:
        print("用法: python validate_docx.py <file.docx> [spec.json]")
        return 1
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"文件不存在: {path}")
        return 1
    spec = None
    if len(sys.argv) >= 3:
        spec_path = Path(sys.argv[2])
        if not spec_path.exists():
            print(f"spec 不存在: {spec_path}")
            return 1
        spec = json.loads(spec_path.read_text(encoding="utf-8-sig"))
    report, problems = validate_docx(path, spec)
    print("=" * 60)
    print(f"验收报告: {path}")
    print("=" * 60)
    for line in report:
        print(line)
    if problems:
        print("\n--- 问题清单 ---")
        for problem in problems:
            print(f"  - {problem}")
        print(f"\n验收未通过（{len(problems)} 项问题），请修复后重新生成/更新。")
        return 1
    print("\n自动 DOCX/XML 检查项全部通过。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
