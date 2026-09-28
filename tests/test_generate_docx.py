# -*- coding: utf-8 -*-
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml.ns import qn


SKILL_DIR = Path(__file__).resolve().parents[1]
GENERATOR = SKILL_DIR / "scripts" / "generate_docx.py"


class GenerateDocxTests(unittest.TestCase):
    def generate(self, spec):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        root = Path(temp_dir.name)
        spec_path = root / "spec.json"
        output_path = root / "result.docx"
        spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(GENERATOR), str(spec_path), str(output_path)],
            capture_output=True,
            text=True,
        )
        return result, output_path

    def base_spec(self, **overrides):
        spec = {
            "cover": False,
            "toc": False,
            "doc_control": False,
            "header_footer": False,
            "body": [{"type": "h1", "text": "测试"}, {"type": "p", "text": "正文内容"}],
        }
        spec.update(overrides)
        return spec

    def test_consulting_profile_controls_body_style_and_page_margins(self):
        result, output = self.generate(self.base_spec(layout_profile="consulting"))
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        self.assertEqual(doc.styles["正文"].font.name, "宋体")
        self.assertAlmostEqual(doc.styles["正文"].font.size.pt, 11.0, places=1)
        self.assertAlmostEqual(doc.sections[0].top_margin.cm, 2.8, places=1)
        self.assertFalse(doc.sections[0].different_first_page_header_footer)
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "")

    def test_figures_are_omitted_by_default_without_error_placeholder(self):
        spec = self.base_spec(body=[
            {"type": "h1", "text": "测试"},
            {"type": "figure", "image": "missing.png", "caption": "不应出现"},
        ])
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        all_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("图片缺失", all_text)
        self.assertNotIn("不应出现", all_text)
        self.assertEqual(len(doc.inline_shapes), 0)

    def test_invalid_table_shape_fails_with_actionable_message(self):
        spec = self.base_spec(body=[
            {"type": "h1", "text": "测试"},
            {"type": "table", "headers": ["A", "B"], "rows": [["only-one"]]},
        ])
        result, _ = self.generate(spec)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("列数", result.stderr + result.stdout)

    def test_cover_and_doc_control_tables_use_stable_pagination_properties(self):
        spec = self.base_spec(
            cover=True,
            doc_control=True,
            project_name="示例项目",
            doc_name="示例文档",
            revision_history=[{"version": "V1.0", "date": "2026-01-01", "content": "初版", "author": "项目组", "reviewer": "", "status": "评审"}],
        )
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        self.assertEqual(len(doc.tables), 2)
        for table in doc.tables:
            layout = table._tbl.tblPr.find(qn("w:tblLayout"))
            self.assertIsNotNone(layout)
            self.assertEqual(layout.get(qn("w:type")), "fixed")
            for row in table.rows:
                tr_pr = row._tr.find(qn("w:trPr"))
                self.assertIsNotNone(tr_pr)
                self.assertIsNotNone(tr_pr.find(qn("w:cantSplit")))

    def test_cover_and_doc_control_tables_use_stable_pagination_properties(self):
        spec = self.base_spec(
            cover=True,
            doc_control=True,
            project_name="示例项目",
            doc_name="示例文档",
            revision_history=[{"version": "V1.0", "date": "2026-01-01", "content": "初版", "author": "项目组", "reviewer": "", "status": "评审"}],
        )
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        self.assertEqual(len(doc.tables), 2)
        for table in doc.tables:
            layout = table._tbl.tblPr.find(qn("w:tblLayout"))
            self.assertIsNotNone(layout)
            self.assertEqual(layout.get(qn("w:type")), "fixed")
            for row in table.rows:
                tr_pr = row._tr.find(qn("w:trPr"))
                self.assertIsNotNone(tr_pr)
                self.assertIsNotNone(tr_pr.find(qn("w:cantSplit")))

    def test_wide_table_uses_landscape_section_and_restores_portrait(self):
        headers = [f"字段{i}" for i in range(8)]
        rows = [[f"值{i}" for i in range(8)]]
        spec = self.base_spec(body=[
            {"type": "h1", "text": "测试"},
            {"type": "table", "headers": headers, "rows": rows},
            {"type": "p", "text": "表后正文"},
        ])
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        self.assertGreaterEqual(len(doc.sections), 3)
        self.assertEqual(doc.sections[-2].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual(doc.sections[-1].orientation, WD_ORIENT.PORTRAIT)

    def test_cover_uses_empty_first_page_header_and_body_header(self):
        spec = self.base_spec(
            cover=True,
            header_footer=True,
            project_name="示例项目",
            doc_name="示例文档",
        )
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        section = doc.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        self.assertEqual(section.first_page_header.paragraphs[0].text, "")
        self.assertEqual(section.first_page_footer.paragraphs[0].text, "")
        self.assertIn("示例项目", section.header.paragraphs[0].text)
        self.assertIn("示例文档", section.header.paragraphs[0].text)
        self.assertIn("第 ", section.footer.paragraphs[0].text)

    def test_header_footer_false_leaves_all_sections_empty(self):
        headers = [f"字段{i}" for i in range(7)]
        spec = self.base_spec(body=[
            {"type": "h1", "text": "测试"},
            {"type": "table", "headers": headers, "rows": [["值"] * 7]},
        ])
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        for section in doc.sections:
            self.assertFalse(section.different_first_page_header_footer)
            self.assertEqual(section.header.paragraphs[0].text, "")
            self.assertEqual(section.footer.paragraphs[0].text, "")

    def test_landscape_sections_link_header_and_footer_to_previous(self):
        headers = [f"字段{i}" for i in range(7)]
        spec = self.base_spec(
            header_footer=True,
            project_name="示例项目",
            doc_name="示例文档",
            body=[
                {"type": "h1", "text": "测试"},
                {"type": "table", "headers": headers, "rows": [["值"] * 7]},
                {"type": "p", "text": "表后正文"},
            ],
        )
        result, output = self.generate(spec)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        doc = Document(output)
        self.assertGreaterEqual(len(doc.sections), 3)
        for section in doc.sections[1:]:
            self.assertTrue(section.header.is_linked_to_previous)
            self.assertTrue(section.footer.is_linked_to_previous)
            self.assertFalse(section.different_first_page_header_footer)


if __name__ == "__main__":
    unittest.main()
